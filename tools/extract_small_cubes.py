#!/usr/bin/env python
"""Extract a two-color double-knitting chart from an xlsx sheet.

Cell colors and values encode the chart:
  cell value 1 (any fill)       -> D  (not a real stitch: the column was
                                       killed by a decrease on the way
                                       to the crown; marks accumulate)
  orange fill (FFFF9900)        -> O  (column does not exist yet: it is
                                       cast on partway up the brim)
  dark blue fill (FF0B5394)     -> B  (blue stitch)
  white / no fill               -> W  (white stitch)

The chart section is a fixed column range (default C..AL). The sheet is
drawn as worn — crown at the top, brim at the bottom — so rows are
emitted in KNIT order: the bottom chart row becomes round 0. Rows with
no B/W cells inside the section (pure margin) are skipped.

Output is a plain CSV: one line per round, comma-separated B/W/O/D.
"""

import argparse
import csv
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

BLUE = {"FF0B5394"}
ORANGE = {"FFFF9900"}
GRAY = {"FFD9EAF7", "FFB7B7B7", "FFCCCCCC", "FFD9D9D9", "FF999999"}


def classify(cell):
    if cell.value in (1, "1"):
        return "D"
    fill = cell.fill
    if fill.fill_type != "solid":
        return "W"
    rgb = str(fill.fgColor.rgb) if fill.fgColor is not None else ""
    if rgb in BLUE:
        return "B"
    if rgb in ORANGE:
        return "O"
    if rgb in GRAY:
        return "D"
    if rgb in ("FFFFFFFF", "00000000"):
        return "W"
    print(f"warning: unrecognized fill {rgb} at {cell.coordinate}, "
          "treating as gray/decreased", file=sys.stderr)
    return "D"


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook")
    ap.add_argument("sheet")
    ap.add_argument("output_csv")
    ap.add_argument("--first-col", default="C")
    ap.add_argument("--last-col", default="AL")
    ap.add_argument("--first-row", type=int, default=2)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook)
    ws = wb[args.sheet]
    c0 = column_index_from_string(args.first_col)
    c1 = column_index_from_string(args.last_col)

    rows = []
    for r in range(args.first_row, ws.max_row + 1):
        symbols = [classify(ws.cell(row=r, column=c)) for c in range(c0, c1 + 1)]
        if any(s in ("B", "W") for s in symbols):
            rows.append(symbols)
    rows.reverse()  # sheet bottom = cast-on = round 0

    with open(args.output_csv, "w", newline="") as f:
        csv.writer(f).writerows(rows)

    width = c1 - c0 + 1
    live = sum(s in ("B", "W") for row in rows for s in row)
    print(f"{args.sheet}: {len(rows)} rounds x {width} columns, "
          f"{live} live stitches per repeat -> {args.output_csv}")


if __name__ == "__main__":
    main()
