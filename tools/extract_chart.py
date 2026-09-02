#!/usr/bin/env python
"""Extract a two-color double-knitting chart from an xlsx sheet.

Cell colors and values encode the chart:
  cell value d (any fill)       -> D  (not a real stitch: the slot was
  cell value 1 (non-blue fill)  -> D   killed by a decrease; the k2tog
                                       goes on its nearest survivor.
                                       Marks may accumulate up the crown
                                       or mark only the diagonal.)
  orange fill (FFFF9900)        -> O  (no stitch in this slot on this
                                       round: not cast on yet, or dead)
  gray fills                    -> D
  dark blue fill (FF0B5394)     -> B  (blue stitch)
  white / no fill               -> W  (white stitch)
Other cell values (design notes such as a, c, 3) are ignored; the fill
decides.

The chart section is a fixed column range (--first-col/--last-col) and
row range (--first-row/--last-row). The sheet is drawn as worn — crown
at the top, brim at the bottom — so rows are emitted in KNIT order: the
bottom chart row becomes round 0. Rows with no B/W cells inside the
section (pure margin) are skipped.

  extract_chart.py patterns/Necker_Birds_Hat.xlsx Small_Cubes out.csv
  extract_chart.py patterns/Necker_Birds_Hat.xlsx Alt_Cubes out.csv \
      --first-col D --last-col AQ --last-row 70

The B/W/O/D symbols are converted with chart.from_legacy into the text
chart format (``f`` / ``b`` / ``f-k2tog`` / ``b-co`` / blank; see
chart.py) and written as CSV, one line per round.
"""

import argparse
import os
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from chart import from_legacy  # noqa: E402

BLUE = {"FF0B5394"}
ORANGE = {"FFFF9900"}
GRAY = {"FFD9EAF7", "FFB7B7B7", "FFCCCCCC", "FFD9D9D9", "FF999999"}


def classify(cell):
    fill = cell.fill
    rgb = (str(fill.fgColor.rgb) if fill.fill_type == "solid"
           and fill.fgColor is not None else "")
    if str(cell.value).strip().lower() == "d":
        return "D"
    if rgb in BLUE:
        return "B"  # digits on a blue cell are planner labels, not marks
    if cell.value in (1, "1"):
        return "D"
    if fill.fill_type != "solid":
        return "W"
    if rgb in ORANGE:
        return "O"
    if rgb in GRAY:
        return "D"
    if rgb in ("FFFFFFFF", "00000000"):
        return "W"
    print(f"warning: unrecognized fill {rgb} at {cell.coordinate}, "
          "treating as gray/decreased", file=sys.stderr)
    return "D"


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook")
    ap.add_argument("sheet")
    ap.add_argument("output_csv")
    ap.add_argument("--first-col", default="C")
    ap.add_argument("--last-col", default="AL")
    ap.add_argument("--first-row", type=int, default=2)
    ap.add_argument("--last-row", type=int, default=None,
                    help="last sheet row of the chart (default: sheet end)")
    ap.add_argument("--implicit-decreases", action="store_true",
                    help="a live slot that turns orange without a d/1 mark "
                         "is still a decrease (k2tog on nearest survivor)")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook)
    ws = wb[args.sheet]
    c0 = column_index_from_string(args.first_col)
    c1 = column_index_from_string(args.last_col)

    rows = []
    last_row = args.last_row or ws.max_row
    for r in range(args.first_row, last_row + 1):
        symbols = [classify(ws.cell(row=r, column=c)) for c in range(c0, c1 + 1)]
        if any(s in ("B", "W") for s in symbols):
            rows.append(symbols)
    rows.reverse()  # sheet bottom = cast-on = round 0

    chart = from_legacy(rows, implicit_decreases=args.implicit_decreases)
    chart.write_csv(args.output_csv)

    width = c1 - c0 + 1
    live = sum(bool(c) for row in chart.rows for c in row)
    print(f"{args.sheet}: {len(rows)} rounds x {width} columns, "
          f"{live} live stitches per repeat -> {args.output_csv}")


if __name__ == "__main__":
    main()
